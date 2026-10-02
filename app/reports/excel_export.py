from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font


def build_attendance_xlsx(title: str, sections: list) -> bytes:
    """One worksheet per section (e.g. "Summary", "Class details")."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, headers, rows in sections:
        sheet = workbook.create_sheet(title=name[:31])
        sheet.append([title])
        sheet["A1"].font = Font(bold=True, size=13)
        sheet.append([])
        sheet.append(headers)
        for cell in sheet[3]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append(row)
        if not rows:
            sheet.append(["No attendance recorded yet."])
        for column in sheet.iter_cols(min_row=3):
            longest = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column)
            sheet.column_dimensions[column[0].column_letter].width = min(longest + 2, 50)
        sheet.freeze_panes = "A4"

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()