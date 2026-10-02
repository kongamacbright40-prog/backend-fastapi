import re
from datetime import datetime, timedelta, timezone

from tests.conftest import PASSWORD, auth


def test_profile_update_and_password_change(client, world):
    student = world["student"]
    me = client.get("/auth/me", headers=student).json()
    assert me["department"]["name"] == "Computer Science"
    assert me["created_at"].endswith("Z")

    r = client.patch("/auth/me", json={"full_name": "Amina B. Bello", "phone_number": "+237600"}, headers=student)
    assert r.status_code == 200 and r.json()["full_name"] == "Amina B. Bello"

    bad = client.post("/auth/change-password", json={"current_password": "nope", "new_password": "NewPass123!"}, headers=student)
    assert bad.status_code == 400
    ok = client.post("/auth/change-password", json={"current_password": PASSWORD, "new_password": "NewPass123!"}, headers=student)
    assert ok.status_code == 204
    auth(client, "student@uni.edu", "NewPass123!")
    client.post(
        "/auth/change-password",
        json={"current_password": "NewPass123!", "new_password": PASSWORD},
        headers=student,
    )


def test_password_reset_flow(client, world, capsys):
    assert client.post("/auth/password-reset/request", json={"email": "lecturer@uni.edu"}).status_code == 204
    code = re.search(r"lecturer@uni.edu: (\d{4})", capsys.readouterr().out).group(1)
    assert client.post("/auth/password-reset/verify", json={"email": "lecturer@uni.edu", "code": "0000" if code != "0000" else "1111"}).status_code == 400
    assert client.post("/auth/password-reset/verify", json={"email": "lecturer@uni.edu", "code": code}).status_code == 204
    r = client.post(
        "/auth/password-reset/confirm",
        json={"email": "lecturer@uni.edu", "code": code, "new_password": PASSWORD},
    )
    assert r.status_code == 204
    # Codes are single-use.
    assert client.post("/auth/password-reset/verify", json={"email": "lecturer@uni.edu", "code": code}).status_code == 400
    # Unknown emails don't leak.
    assert client.post("/auth/password-reset/request", json={"email": "ghost@uni.edu"}).status_code == 204


def test_password_reset_code_is_burned_after_too_many_wrong_guesses(client, world, capsys):
    from app.auth.service import MAX_RESET_ATTEMPTS

    assert client.post("/auth/password-reset/request", json={"email": "lecturer@uni.edu"}).status_code == 204
    code = re.search(r"lecturer@uni.edu: (\d{4})", capsys.readouterr().out).group(1)
    wrong = "0000" if code != "0000" else "1111"
    for _ in range(MAX_RESET_ATTEMPTS):
        r = client.post("/auth/password-reset/verify", json={"email": "lecturer@uni.edu", "code": wrong})
        assert r.status_code == 400
    # The right code no longer works once the attempts are used up.
    r = client.post(
        "/auth/password-reset/confirm",
        json={"email": "lecturer@uni.edu", "code": code, "new_password": "Hacked123!"},
    )
    assert r.status_code == 400
    auth(client, "lecturer@uni.edu")


def test_admin_registration_requires_the_code_once_an_admin_exists(client, world, monkeypatch):
    from app.config import settings

    body = {
        "full_name": "Second Admin",
        "email": "admin2@uni.edu",
        "phone_number": "+237600000000",
        "matricule_number": "ADM-002",
        "password": PASSWORD,
        "role": "admin",
    }
    monkeypatch.setattr(settings, "admin_registration_code", "")
    closed = client.post("/auth/register", json=body)
    assert closed.status_code == 403 and "closed" in closed.json()["detail"]

    monkeypatch.setattr(settings, "admin_registration_code", "S3CRET")
    assert client.post("/auth/register", json={**body, "admin_code": "nope"}).status_code == 403
    ok = client.post("/auth/register", json={**body, "admin_code": "S3CRET"})
    assert ok.status_code == 201, ok.text
    created = ok.json()
    assert created["role"] == "admin"
    assert created["phone_number"] == "+237600000000"
    assert created["matricule_number"] == "ADM-002"
    auth(client, "admin2@uni.edu")

    # An admin added by another admin (no password) activates without the code.
    r = client.post(
        "/admin/users",
        json={"full_name": "Invited Admin", "email": "admin3@uni.edu", "role": "admin"},
        headers=world["admin"],
    )
    assert r.status_code == 201, r.text
    activated = client.post(
        "/auth/register",
        json={"full_name": "Invited Admin", "email": "admin3@uni.edu", "password": PASSWORD, "role": "admin"},
    )
    assert activated.status_code == 201, activated.text
    auth(client, "admin3@uni.edu")


def test_short_matricule_and_short_password_are_accepted(client, world):
    body = {"full_name": "Short Id", "email": "short.id@uni.edu", "role": "lecturer", "matricule_number": "L2"}
    assert client.post("/auth/register", json={**body, "password": "abc"}).status_code == 422
    r = client.post("/auth/register", json={**body, "password": "abcd"})
    assert r.status_code == 201, r.text
    assert r.json()["matricule_number"] == "L2"
    auth(client, "short.id@uni.edu", "abcd")


def test_login_upgrades_slow_password_hashes(client, world):
    import bcrypt

    from app.auth.models import Profile
    from app.auth.security import BCRYPT_ROUNDS
    from app.database import SessionLocal

    client.post(
        "/auth/register",
        json={"full_name": "Old Hash", "email": "old.hash@uni.edu", "password": PASSWORD, "role": "lecturer", "matricule_number": "OLD-1"},
    )
    with SessionLocal() as db:
        user = db.query(Profile).filter(Profile.email == "old.hash@uni.edu").one()
        user.hashed_password = bcrypt.hashpw(PASSWORD.encode(), bcrypt.gensalt(12)).decode()
        db.commit()
    auth(client, "old.hash@uni.edu")
    with SessionLocal() as db:
        stored = db.query(Profile).filter(Profile.email == "old.hash@uni.edu").one().hashed_password
    assert stored.split("$")[2] == f"{BCRYPT_ROUNDS:02d}"
    auth(client, "old.hash@uni.edu")


def test_email_is_case_and_space_insensitive(client, world, capsys):
    r = client.post("/auth/login", json={"email": "  Student@Uni.EDU ", "password": PASSWORD})
    assert r.status_code == 200, r.text
    # Registering the same address with different capitals is a duplicate.
    dup = client.post(
        "/auth/register",
        json={"full_name": "Dup", "email": "STUDENT@uni.edu", "password": PASSWORD, "role": "student", "matricule_number": "DUP-1"},
    )
    assert dup.status_code == 400
    assert client.post("/auth/password-reset/request", json={"email": "Student@Uni.edu"}).status_code == 204
    assert "student@uni.edu" in capsys.readouterr().out


def test_refresh_rejects_bad_tokens_with_401(client, world):
    from jose import jwt

    from app.config import settings

    login = client.post("/auth/login", json={"email": "student@uni.edu", "password": PASSWORD}).json()
    ok = client.post("/auth/refresh", json={"refresh_token": login["refresh_token"]})
    assert ok.status_code == 200 and ok.json()["access_token"]

    # An access token is not a refresh token.
    assert client.post("/auth/refresh", json={"refresh_token": login["access_token"]}).status_code == 401
    assert client.post("/auth/refresh", json={"refresh_token": "garbage"}).status_code == 401
    for claims in ({"type": "refresh"}, {"type": "refresh", "sub": "abc"}):
        token = jwt.encode(claims, settings.jwt_secret_key, algorithm=settings.jwt_algorithm)
        assert client.post("/auth/refresh", json={"refresh_token": token}).status_code == 401


def test_admin_created_account_is_activated_by_owner(client, world):
    admin = world["admin"]
    r = client.post(
        "/admin/users",
        json={"full_name": "Chidi Okafor", "email": "chidi@uni.edu", "role": "student", "department_id": world["dept"]["id"]},
        headers=admin,
    )
    assert r.status_code == 201, r.text
    assert r.json()["pending_activation"] is True
    assert client.post("/auth/login", json={"email": "chidi@uni.edu", "password": "whatever1"}).status_code == 403

    activated = client.post(
        "/auth/register",
        json={
            "full_name": "Chidi",
            "email": "chidi@uni.edu",
            "password": PASSWORD,
            "role": "student",
            "matricule_number": "ICT-2",
        },
    )
    assert activated.status_code == 201, activated.text
    body = activated.json()
    assert body["pending_activation"] is False
    assert body["full_name"] == "Chidi Okafor"  # admin-provided name is kept
    assert body["matricule_number"] == "ICT-2"
    auth(client, "chidi@uni.edu")


def test_admin_user_management(client, world):
    admin = world["admin"]
    created = client.post(
        "/admin/users",
        json={"full_name": "Temp User", "email": "temp@uni.edu", "role": "lecturer", "password": PASSWORD, "matricule_number": "FAC-9"},
        headers=admin,
    ).json()
    uid = created["id"]
    r = client.patch(f"/admin/users/{uid}", json={"full_name": "Temp Lecturer", "phone_number": "123"}, headers=admin)
    assert r.json()["full_name"] == "Temp Lecturer"

    assert client.patch(f"/admin/users/{uid}/status", json={"is_active": False}, headers=admin).json()["is_active"] is False
    assert client.post("/auth/login", json={"email": "temp@uni.edu", "password": PASSWORD}).status_code == 403
    client.patch(f"/admin/users/{uid}/status", json={"is_active": True}, headers=admin)

    found = client.get("/admin/users", params={"q": "temp"}, headers=admin).json()
    assert [u["id"] for u in found["items"]] == [uid]

    assert client.delete(f"/admin/users/{uid}", headers=admin).status_code == 204
    assert client.get("/admin/users", params={"q": "temp"}, headers=admin).json()["total"] == 0

    activity = client.get("/admin/activity", headers=admin).json()
    assert any(a["title"] == "User deleted" for a in activity)


def test_course_management_and_enrollment(client, world):
    admin, lecturer, student = world["admin"], world["lecturer"], world["student"]
    course_id = world["course"]["id"]

    # Student is enrolled through their department.
    mine = client.get("/courses/mine", headers=student).json()
    assert [c["code"] for c in mine["items"]] == ["CS-301"]
    assert client.get("/courses/mine", headers=lecturer).json()["total"] == 0

    r = client.post(f"/courses/{course_id}/assign-lecturer", json={"lecturer_id": world["lecturer_id"]}, headers=admin)
    assert r.json()["lecturer"]["full_name"] == "Kofi Mensah"
    assert [c["id"] for c in client.get("/courses/mine", headers=lecturer).json()["items"]] == [course_id]

    patched = client.patch(f"/courses/{course_id}", json={"description": "Trees and graphs"}, headers=admin).json()
    assert patched["description"] == "Trees and graphs" and patched["credits"] == 4

    roster = client.get(f"/courses/{course_id}/roster", headers=lecturer).json()
    assert {s["full_name"] for s in roster} >= {"Amina B. Bello"}

    # Explicit enrolment for a student from another department.
    other = client.post(
        "/admin/users",
        json={"full_name": "Visiting Student", "email": "visit@uni.edu", "role": "student", "password": PASSWORD, "matricule_number": "VIS-1"},
        headers=admin,
    ).json()
    assert client.get("/courses/mine", headers=auth(client, "visit@uni.edu")).json()["total"] == 0
    client.post(f"/courses/{course_id}/enrollments", json={"profile_id": other["id"]}, headers=admin)
    assert client.get("/courses/mine", headers=auth(client, "visit@uni.edu")).json()["total"] == 1

    notes = client.get("/notifications", headers=world["lecturer"]).json()["items"]
    assert any(n["type"] == "new_course" for n in notes)


def test_departments_update_and_archive(client, world):
    admin = world["admin"]
    dept = client.post("/departments", json={"name": "Physics", "faculty_id": world["faculty"]["id"]}, headers=admin).json()
    renamed = client.patch(f"/departments/{dept['id']}", json={"name": "Applied Physics"}, headers=admin).json()
    assert renamed["name"] == "Applied Physics"
    archived = client.post(f"/departments/{dept['id']}/archive", headers=admin).json()
    assert archived["is_active"] is False
    faculties = client.get("/faculties", headers=admin).json()["items"]
    assert faculties[0]["department_count"] == 2


def test_schedule_start_participate_and_report(client, world):
    admin, lecturer, student = world["admin"], world["lecturer"], world["student"]
    course_id = world["course"]["id"]
    start = (datetime.now(timezone.utc) + timedelta(hours=1)).isoformat()

    scheduled = client.post(
        "/classes/schedule",
        json={"course_id": course_id, "title": "Graphs", "scheduled_start": start, "duration_minutes": 90},
        headers=lecturer,
    )
    assert scheduled.status_code == 201, scheduled.text
    session = scheduled.json()
    sid = session["id"]
    assert session["status"] == "scheduled" and session["expected_count"] >= 1

    student_classes = client.get("/classes", headers=student).json()["items"]
    assert sid in [c["id"] for c in student_classes]
    assert any(n["type"] == "class_reminder" for n in client.get("/notifications", headers=student).json()["items"])

    # Students can't join before the class starts.
    student_token = student["Authorization"].split()[1]
    lecturer_token = lecturer["Authorization"].split()[1]
    try:
        with client.websocket_connect(f"/ws/signal/{sid}?token={student_token}") as ws:
            ws.receive_json()
        joined_early = True
    except Exception:
        joined_early = False
    assert not joined_early

    live = client.post(f"/classes/{sid}/start", headers=lecturer).json()
    assert live["status"] == "live"
    assert [c["id"] for c in client.get("/classes", params={"live_only": True}, headers=student).json()["items"]] == [sid]

    with client.websocket_connect(f"/ws/signal/{sid}?token={lecturer_token}") as lws:
        assert lws.receive_json()["type"] == "room_state"
        with client.websocket_connect(f"/ws/signal/{sid}?token={student_token}") as sws:
            assert sws.receive_json()["type"] == "room_state"
            assert lws.receive_json()["type"] == "peer_joined"

            participants = client.get(f"/classes/{sid}/participants", headers=lecturer).json()
            assert {p["name"] for p in participants} == {"Kofi Mensah", "Amina B. Bello"}

            hands = client.put(f"/classes/{sid}/hand", json={"raised": True}, headers=student).json()
            assert [p["is_hand_raised"] for p in hands if p["name"] == "Amina B. Bello"] == [True]

            msg = client.post(f"/classes/{sid}/messages", json={"message": "Hello!", "is_question": True}, headers=student)
            assert msg.status_code == 201
            messages = client.get(f"/classes/{sid}/messages", headers=lecturer).json()
            assert messages[0]["sender_name"] == "Amina B. Bello" and messages[0]["is_question"] is True
            assert client.get(f"/classes/{sid}/messages", params={"after_id": messages[0]["id"]}, headers=lecturer).json() == []

            question = client.post(
                f"/participation/sessions/{sid}/questions",
                json={"prompt": "BFS uses?", "options": [{"text": "Queue", "is_correct": True}, {"text": "Stack"}]},
                headers=lecturer,
            ).json()
            answer = client.post(
                f"/participation/questions/{question['id']}/respond",
                json={"selected_option_id": question["options"][0]["id"]},
                headers=student,
            )
            assert answer.json()["is_correct"] is True
            mine_answer = client.get(f"/participation/questions/{question['id']}/responses/me", headers=student)
            assert mine_answer.json()["selected_option_id"] == question["options"][0]["id"]
            listed = client.get(f"/participation/sessions/{sid}/questions", headers=lecturer).json()["items"]
            assert [o["response_count"] for o in listed[0]["options"]] == [1, 0]

            draft = client.post(
                f"/participation/sessions/{sid}/questions",
                json={"prompt": "DFS uses?", "options": [{"text": "Stack", "is_correct": True}, {"text": "Queue"}], "launch": False},
                headers=lecturer,
            ).json()
            assert draft["is_open"] is False
            assert client.post(f"/participation/questions/{draft['id']}/launch", headers=lecturer).json()["is_open"] is True
            open_now = client.get(f"/participation/sessions/{sid}/questions/open", headers=student).json()["items"]
            assert [q["id"] for q in open_now] == [draft["id"]]

    ended = client.post(f"/classes/{sid}/end", headers=lecturer).json()
    assert ended["status"] == "completed"

    mine = client.get("/attendance/me", headers=student).json()
    assert mine[0]["session"]["id"] == sid and mine[0]["status"] == "present"
    assert mine[0]["session"]["course"]["code"] == "CS-301"

    summary = client.get(f"/attendance/courses/{course_id}/summary", headers=lecturer).json()
    amina = next(s for s in summary if s["profile"]["full_name"] == "Amina B. Bello")
    assert (amina["total_sessions"], amina["present_count"], amina["absent_count"]) == (1, 1, 0)

    appeal = client.post(f"/attendance/sessions/{sid}/appeals", json={"reason": "Network dropped"}, headers=student)
    assert appeal.status_code == 201
    assert client.post(f"/attendance/sessions/{sid}/appeals", json={"reason": "Again"}, headers=student).status_code == 409
    assert any(n["title"] == "Attendance appeal" for n in client.get("/notifications", headers=lecturer).json()["items"])

    report = client.get("/reports/lecturer", headers=lecturer).json()
    assert report["metrics"]["sessions"] == 1
    assert report["breakdown"][0]["label"] == "CS-301"
    overview = client.get("/reports/overview", headers=admin).json()
    assert overview["metrics"]["total_students"] >= 1 and "avg_attendance" in overview["metrics"]
    institution = client.get("/reports/institution", headers=admin).json()
    assert {b["meta"]["type"] for b in institution["breakdown"]} == {"faculty", "course"}

    for params in ({"session_id": sid}, {"course_id": course_id}, {}):
        pdf = client.get("/reports/attendance.pdf", params=params, headers=lecturer)
        assert pdf.status_code == 200 and pdf.headers["content-type"] == "application/pdf"
    xlsx = client.get("/reports/attendance.xlsx", params={"course_id": course_id, "utc_offset_minutes": 60}, headers=admin)
    assert xlsx.status_code == 200
    from io import BytesIO

    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(xlsx.content))
    assert workbook.sheetnames == ["Summary", "Class details"]
    details = workbook["Class details"]
    headers = [c.value for c in details[3]]
    assert headers[5] == "Joined (UTC+01:00)" and headers[6] == "Left (UTC+01:00)"
    rows = [[c.value for c in row] for row in details.iter_rows(min_row=4)]
    student_row = next(r for r in rows if r[2] == "Amina B. Bello")
    assert re.fullmatch(r"\d\d:\d\d:\d\d", student_row[5]) and re.fullmatch(r"\d\d:\d\d:\d\d", student_row[6])
    summary_headers = [c.value for c in workbook["Summary"][3]]
    assert summary_headers == ["Matricule", "Full name", "Role", "Attendance", "Total minutes"]
    csv = client.get("/reports/attendance.csv", params={"session_id": sid}, headers=lecturer)
    assert csv.status_code == 200 and "Amina B. Bello" in csv.content.decode("utf-8-sig")
    course_csv = client.get("/reports/attendance.csv", params={"course_id": course_id}, headers=lecturer).content.decode("utf-8-sig")
    assert "Class details" in course_csv and "Joined (UTC+00:00)" in course_csv

    excused = client.post(
        f"/attendance/sessions/{sid}/mark",
        json={"profile_id": world["student_id"], "status": "excused"},
        headers=lecturer,
    )
    assert excused.status_code == 200, excused.text
    summary = client.get(f"/attendance/courses/{course_id}/summary", headers=lecturer).json()
    amina = next(s for s in summary if s["profile"]["full_name"] == "Amina B. Bello")
    assert (amina["present_count"], amina["excused_count"], amina["absent_count"]) == (0, 1, 0)


def test_notifications_read_state(client, world):
    student = world["student"]
    items = client.get("/notifications", headers=student).json()["items"]
    assert items
    first = client.post(f"/notifications/{items[0]['id']}/read", headers=student).json()
    assert first["is_read"] is True
    assert client.post("/notifications/read-all", headers=student).status_code == 204
    assert client.get("/notifications", params={"unread_only": True}, headers=student).json()["total"] == 0


def test_terms_and_settings(client, world):
    admin, student = world["admin"], world["student"]
    now = datetime.now(timezone.utc)
    term = {
        "name": "Fall Semester 2026",
        "code": "FA26",
        "academic_year": "2026/2027",
        "start_date": (now - timedelta(days=30)).isoformat(),
        "end_date": (now + timedelta(days=60)).isoformat(),
        "status": "active",
        "enrollment_open": True,
    }
    created = client.post("/academic-terms", json=term, headers=admin)
    assert created.status_code == 201, created.text
    body = created.json()
    assert body["course_count"] >= 1 and body["start_date"].endswith("Z")
    assert client.post("/academic-terms", json=term, headers=admin).status_code == 400  # duplicate code
    updated = client.put(f"/academic-terms/{body['id']}", json={**term, "notes": "Exams in week 14"}, headers=admin)
    assert updated.json()["notes"] == "Exams in week 14"
    assert len(client.get("/academic-terms", headers=student).json()) == 1

    settings = client.get("/settings/system", headers=student).json()
    assert settings["minimum_attendance"] == 75.0
    new = {**settings, "minimum_attendance": 80, "late_threshold_minutes": 10}
    new.pop("cluster_version"), new.pop("last_synced_at")
    assert client.put("/settings/system", json=new, headers=student).status_code == 403
    saved = client.put("/settings/system", json=new, headers=admin).json()
    assert saved["minimum_attendance"] == 80 and saved["last_synced_at"].endswith("Z")


def test_stale_live_classes_are_closed_and_admin_can_end_a_class(client, world):
    from datetime import datetime, timedelta

    from app.classes.models import ClassSession
    from app.database import SessionLocal

    lecturer, admin = world["lecturer"], world["admin"]
    course = client.post(
        "/courses",
        json={"code": "STALE-101", "title": "Stale Classes", "department_id": world["dept"]["id"]},
        headers=admin,
    ).json()
    stale = client.post("/classes", json={"course_id": course["id"], "title": "Forgotten"}, headers=lecturer).json()
    with SessionLocal() as db:
        db.get(ClassSession, stale["id"]).started_at = datetime.utcnow() - timedelta(days=2)
        db.commit()
    live_ids = [s["id"] for s in client.get("/admin/sessions?live_only=true", headers=admin).json()["items"]]
    assert stale["id"] not in live_ids
    assert client.get(f"/classes/{stale['id']}", headers=admin).json()["status"] == "completed"

    fresh = client.post("/classes", json={"course_id": course["id"], "title": "Running"}, headers=lecturer).json()
    overview = client.get("/reports/overview", headers=admin).json()
    assert overview["metrics"]["active_classes"] >= 1
    assert client.post(f"/classes/{fresh['id']}/end", headers=world["student"]).status_code == 403
    ended = client.post(f"/classes/{fresh['id']}/end", headers=admin)
    assert ended.status_code == 200 and ended.json()["status"] == "completed"


def test_registration_departments_and_student_self_enrollment(client, world):
    admin, student = world["admin"], world["student"]
    public = client.get("/departments/public")
    assert public.status_code == 200
    assert {"id": world["dept"]["id"], "name": "Computer Science"}.items() <= public.json()[0].items()

    other_dept = client.post("/departments", json={"name": "Mathematics", "faculty_id": world["faculty"]["id"]}, headers=admin).json()
    math = client.post("/courses", json={"code": "MTH-101", "title": "Calculus", "department_id": other_dept["id"]}, headers=admin).json()

    reg = client.post(
        "/auth/register",
        json={"full_name": "Dept Student", "email": "dept.student@uni.edu", "password": PASSWORD, "role": "student", "matricule_number": "DS-1", "department_id": other_dept["id"]},
    )
    assert reg.status_code == 201 and reg.json()["department"]["id"] == other_dept["id"]

    listed = {c["code"]: c for c in client.get("/courses", headers=student).json()["items"]}
    assert listed["CS-301"]["enrollment"] == "department"
    assert listed["MTH-101"]["enrollment"] is None
    assert client.get("/courses", headers=admin).json()["items"][0]["enrollment"] is None

    joined = client.post(f"/courses/{math['id']}/enroll", headers=student)
    assert joined.status_code == 200 and joined.json()["enrollment"] == "explicit"
    mine = [c["code"] for c in client.get("/courses/mine", headers=student).json()["items"]]
    assert "MTH-101" in mine
    assert client.post(f"/courses/{math['id']}/enroll", headers=admin).status_code == 403

    assert client.delete(f"/courses/{world['course']['id']}/enroll", headers=student).status_code == 400
    dropped = client.delete(f"/courses/{math['id']}/enroll", headers=student)
    assert dropped.status_code == 200 and dropped.json()["enrollment"] is None
    assert "MTH-101" not in [c["code"] for c in client.get("/courses/mine", headers=student).json()["items"]]


def test_whiteboard_is_relayed_to_students_and_replayed_to_late_joiners(client, world):
    admin, lecturer, student = world["admin"], world["lecturer"], world["student"]
    course = client.post(
        "/courses",
        json={"code": "BOARD-1", "title": "Board Course", "department_id": world["dept"]["id"]},
        headers=admin,
    ).json()
    sid = client.post("/classes", json={"course_id": course["id"], "title": "Board"}, headers=lecturer).json()["id"]
    lecturer_token = lecturer["Authorization"].split()[1]
    student_token = student["Authorization"].split()[1]

    with client.websocket_connect(f"/ws/signal/{sid}?token={lecturer_token}") as lws:
        assert lws.receive_json()["type"] == "room_state"
        with client.websocket_connect(f"/ws/signal/{sid}?token={student_token}") as sws:
            assert sws.receive_json()["type"] == "room_state"
            assert lws.receive_json()["type"] == "peer_joined"

            lws.send_json({"type": "board", "op": "show"})
            assert sws.receive_json() == {"type": "board", "op": "show"}
            lws.send_json({"type": "board", "op": "begin", "id": "s1", "color": 4280391411, "width": 0.01, "points": [[0.1, 0.2]]})
            begin = sws.receive_json()
            assert begin["op"] == "begin" and begin["points"] == [[0.1, 0.2]]
            lws.send_json({"type": "board", "op": "extend", "id": "s1", "points": [[0.3, 0.4], [0.5, 0.6]]})
            assert sws.receive_json()["points"] == [[0.3, 0.4], [0.5, 0.6]]

            # Students can't draw, malformed operations are dropped.
            sws.send_json({"type": "board", "op": "clear"})
            lws.send_json({"type": "board", "op": "begin", "id": "bad", "color": "red", "width": 1, "points": []})
            lws.send_json({"type": "board", "op": "begin", "id": "s2", "color": 1, "width": 0.02, "points": [[0.9, 0.9]]})
            assert sws.receive_json()["id"] == "s2"

            # Screen share on/off is announced to the class.
            lws.send_json({"type": "board", "op": "screen", "on": True})
            assert sws.receive_json() == {"type": "board", "op": "screen", "on": True}

        # A student joining later gets the whole board (and the share state).
        with client.websocket_connect(f"/ws/signal/{sid}?token={student_token}") as late:
            assert late.receive_json()["type"] == "room_state"
            state = late.receive_json()
            assert state["type"] == "board_state" and state["active"] is True
            assert state["screen_on"] is True
            assert [s["id"] for s in state["strokes"]] == ["s1", "s2"]
            assert state["strokes"][0]["points"] == [[0.1, 0.2], [0.3, 0.4], [0.5, 0.6]]

    client.post(f"/classes/{sid}/end", headers=lecturer)
    from app.signaling.connection_manager import manager

    assert sid not in manager.boards


def test_auto_join_leave_recording_can_be_turned_off(client, world):
    admin, lecturer, student = world["admin"], world["lecturer"], world["student"]
    settings = client.get("/settings/system", headers=admin).json()
    body = {k: v for k, v in settings.items() if k not in ("id", "updated_at", "updated_by")}
    client.put("/settings/system", json={**body, "auto_join_leave_recording": False}, headers=admin)
    course = client.post(
        "/courses",
        json={"code": "NOREC-1", "title": "No Recording", "department_id": world["dept"]["id"]},
        headers=admin,
    ).json()
    sid = client.post("/classes", json={"course_id": course["id"]}, headers=lecturer).json()["id"]
    token = student["Authorization"].split()[1]
    try:
        with client.websocket_connect(f"/ws/signal/{sid}?token={token}") as ws:
            assert ws.receive_json()["type"] == "room_state"
        records = client.get(f"/attendance/sessions/{sid}", headers=lecturer).json()["items"]
        assert records == []
    finally:
        client.put("/settings/system", json={**body, "auto_join_leave_recording": True}, headers=admin)


def test_class_everyone_left_is_ended_after_15_minutes(client, world):
    from datetime import datetime, timedelta

    from app.signaling.connection_manager import manager

    admin, lecturer = world["admin"], world["lecturer"]
    course = client.post(
        "/courses",
        json={"code": "ABAND-1", "title": "Abandoned", "department_id": world["dept"]["id"]},
        headers=admin,
    ).json()
    sid = client.post("/classes", json={"course_id": course["id"], "title": "Left open"}, headers=lecturer).json()["id"]
    token = lecturer["Authorization"].split()[1]
    with client.websocket_connect(f"/ws/signal/{sid}?token={token}") as ws:
        assert ws.receive_json()["type"] == "room_state"

    # Lecturer just left: the class is still live (they may come back).
    assert sid in manager.emptied_at
    live = [s["id"] for s in client.get("/admin/sessions?live_only=true", headers=admin).json()["items"]]
    assert sid in live

    left = datetime.utcnow() - timedelta(minutes=20)
    from app.classes.models import ClassSession
    from app.database import SessionLocal

    with SessionLocal() as db:
        db.get(ClassSession, sid).started_at = datetime.utcnow() - timedelta(minutes=40)
        db.commit()
    manager.emptied_at[sid] = left
    live = [s["id"] for s in client.get("/admin/sessions?live_only=true", headers=admin).json()["items"]]
    assert sid not in live
    ended = client.get(f"/classes/{sid}", headers=admin).json()
    assert ended["status"] == "completed"
    assert ended["ended_at"].startswith(left.strftime("%Y-%m-%dT%H:%M"))


def test_cors_allows_deployed_web_app_and_localhost_only(client):
    def allowed(origin):
        r = client.options("/auth/login", headers={"Origin": origin, "Access-Control-Request-Method": "POST"})
        return r.headers.get("access-control-allow-origin") == origin

    assert allowed("https://smartclass.example.edu")
    assert allowed("https://admin.example.edu")  # trailing slash in the setting is ignored
    assert allowed("http://localhost:8080")
    assert not allowed("https://evil.example.com")
